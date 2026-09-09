import json
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

with open("/tmp/shuffled_50.json", encoding="utf-8") as f:
    questions = json.load(f)

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

def add_run(p, text, bold=False, size=12):
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.bold = bold
    return r

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p, "KUNCI JAWABAN", bold=True, size=16)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p, "Truth Class — Track SBY — Roma-Korintus (Message 7-9)", bold=True, size=12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p, "Untuk pengoreksi (ID)", bold=False, size=11)

doc.add_paragraph()

# Grid: 5 blocks of 10 rows, each block = No|Jwb columns -> 10 columns total, 10 rows
n_blocks = 5
rows_per_block = 10
table = doc.add_table(rows=rows_per_block + 1, cols=n_blocks * 2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# header row
for b in range(n_blocks):
    c1 = table.cell(0, b*2)
    c2 = table.cell(0, b*2+1)
    c1.text = "No"
    c2.text = "Jwb"
    for cell in (c1, c2):
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in para.runs:
                r.bold = True

answers = {i+1: q["correct_letter"] for i, q in enumerate(questions)}

for b in range(n_blocks):
    for row in range(rows_per_block):
        qno = b * rows_per_block + row + 1
        c1 = table.cell(row+1, b*2)
        c2 = table.cell(row+1, b*2+1)
        c1.text = str(qno)
        c2.text = answers[qno]
        for cell in (c1, c2):
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER

# borders
def set_table_borders(table):
    from docx.oxml import OxmlElement
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '1F3864')
        borders.append(el)
    tblPr.append(borders)

set_table_borders(table)

doc.save("/tmp/Kunci_Jawaban_Truth_M07-09_SBY.docx")
print("saved key")
print("Distribution check:", {l: list(answers.values()).count(l) for l in "ABCD"})
